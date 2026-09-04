# -*- coding: utf-8 -*-
"""Export deliverables Markdown to DOCX / PPTX without changing the .md sources."""
from __future__ import print_function

import re
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation
from pptx.util import Inches as PInches, Pt as PPt
from pptx.dml.color import RGBColor as PRgb
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

HERE = Path(__file__).resolve().parent
OUT = HERE / "office"
BG = PRgb(0x0F, 0x1F, 0x1A)
GOLD = PRgb(0xC4, 0xB5, 0x81)
TEXT = PRgb(0xF3, 0xF0, 0xE6)
MUTED = PRgb(0xB8, 0xC4, 0xBC)


def strip_md_links(s):
    s = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", s)
    s = s.replace("&lt;", "<").replace("&gt;", ">").replace("&amp;", "&")
    return s


def add_runs(paragraph, text, bold_default=False, size=11):
    text = strip_md_links(text)
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
        else:
            run = paragraph.add_run(part)
            if bold_default:
                run.bold = True
        run.font.size = Pt(size)


def md_to_docx(md_path, docx_path):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    in_code = False
    code_buf = []
    # skip marp front matter
    if lines and lines[0].strip() == "---":
        i = 1
        while i < len(lines) and lines[i].strip() != "---":
            i += 1
        if i < len(lines):
            i += 1

    while i < len(lines):
        line = lines[i]
        raw = line.rstrip("\n")

        if raw.strip().startswith("```"):
            if not in_code:
                in_code = True
                code_buf = []
            else:
                in_code = False
                p = doc.add_paragraph()
                run = p.add_run("\n".join(code_buf))
                run.font.name = "Consolas"
                run.font.size = Pt(9)
            i += 1
            continue

        if in_code:
            code_buf.append(raw)
            i += 1
            continue

        if raw.strip() == "---" or raw.strip().startswith("<!--"):
            i += 1
            continue

        if raw.startswith("# "):
            doc.add_heading(strip_md_links(raw[2:].strip()), level=0)
            i += 1
            continue
        if raw.startswith("## "):
            doc.add_heading(strip_md_links(raw[3:].strip()), level=1)
            i += 1
            continue
        if raw.startswith("### "):
            doc.add_heading(strip_md_links(raw[4:].strip()), level=2)
            i += 1
            continue

        if raw.strip().startswith("|") and i + 1 < len(lines) and re.match(r"^\|?\s*[-:| ]+\|", lines[i + 1]):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not re.match(r"^[-: ]+$", "".join(row).replace("|", "")):
                    # skip separator by checking cells
                    if not all(re.match(r"^:?-+:?$", c or "-") for c in row):
                        rows.append(row)
                i += 1
            if rows:
                cols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=cols)
                table.style = "Table Grid"
                for r_idx, row in enumerate(rows):
                    for c_idx in range(cols):
                        cell = table.rows[r_idx].cells[c_idx]
                        cell.text = ""
                        p = cell.paragraphs[0]
                        add_runs(p, row[c_idx] if c_idx < len(row) else "", size=10)
                        if r_idx == 0:
                            for run in p.runs:
                                run.bold = True
            continue

        m = re.match(r"^(\s*)([-*+]|\d+\.)\s+(.*)$", raw)
        if m:
            body = m.group(3)
            style_name = "List Number" if m.group(2)[-1] == "." else "List Bullet"
            try:
                p = doc.add_paragraph(style=style_name)
            except KeyError:
                p = doc.add_paragraph(style="List Bullet")
            add_runs(p, body)
            i += 1
            continue

        if not raw.strip():
            i += 1
            continue

        p = doc.add_paragraph()
        add_runs(p, raw.strip())
        i += 1

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))
    print("DOCX", docx_path.name)


def set_slide_bg(slide, rgb):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb


def add_textbox(slide, left, top, width, height, text, size=20, bold=False, color=TEXT, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = PPt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"
    return tf


def md_slides_to_pptx(md_path, pptx_path):
    text = md_path.read_text(encoding="utf-8")
    # drop marp front matter
    if text.startswith("---"):
        end = text.find("\n---", 3)
        if end != -1:
            text = text[end + 4 :]

    chunks = re.split(r"\n---\n", text)
    prs = Presentation()
    prs.slide_width = PInches(13.333)
    prs.slide_height = PInches(7.5)
    blank = prs.slide_layouts[6]

    for chunk in chunks:
        lines = [ln for ln in chunk.splitlines() if not ln.strip().startswith("<!--")]
        # trim empties
        while lines and not lines[0].strip():
            lines.pop(0)
        while lines and not lines[-1].strip():
            lines.pop()
        if not lines:
            continue

        slide = prs.slides.add_slide(blank)
        set_slide_bg(slide, BG)

        title = None
        subtitle = None
        body_lines = []
        idx = 0
        if lines[0].startswith("# "):
            title = strip_md_links(lines[0][2:].strip())
            idx = 1
            if idx < len(lines) and lines[idx].startswith("## "):
                subtitle = strip_md_links(lines[idx][3:].strip())
                idx += 1
        body_lines = lines[idx:]

        y = PInches(0.4)
        if title:
            add_textbox(slide, PInches(0.6), y, PInches(12), PInches(0.9), title, size=32, bold=True, color=TEXT)
            y = PInches(1.2)
            # gold bar
            shape = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, PInches(0.6), PInches(1.15), PInches(1.2), PInches(0.06)
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = GOLD
            shape.line.fill.background()
        if subtitle:
            add_textbox(slide, PInches(0.6), y, PInches(12), PInches(0.5), subtitle, size=20, bold=False, color=GOLD)
            y = PInches(1.7)

        # flatten body into text with simple formatting
        body = []
        in_code = False
        for ln in body_lines:
            if ln.strip().startswith("```"):
                in_code = not in_code
                continue
            if in_code:
                body.append(ln)
                continue
            if ln.strip().startswith("|"):
                # table row as plain text
                cells = [c.strip() for c in ln.strip().strip("|").split("|")]
                if all(re.match(r"^:?-+:?$", c or "-") for c in cells):
                    continue
                body.append("  |  ".join(strip_md_links(c) for c in cells))
                continue
            if not ln.strip():
                body.append("")
                continue
            body.append(strip_md_links(ln.rstrip()))

        body_text = "\n".join(body).strip()
        # clean markdown bold/code markers lightly for slides
        body_text = re.sub(r"\*\*([^*]+)\*\*", r"\1", body_text)
        body_text = re.sub(r"`([^`]+)`", r"\1", body_text)

        if body_text:
            box = slide.shapes.add_textbox(PInches(0.6), y, PInches(12.1), PInches(5.6))
            tf = box.text_frame
            tf.word_wrap = True
            first = True
            for para_line in body_text.split("\n"):
                if first:
                    p = tf.paragraphs[0]
                    first = False
                else:
                    p = tf.add_paragraph()
                p.alignment = PP_ALIGN.LEFT
                run = p.add_run()
                run.text = para_line if para_line else " "
                run.font.size = PPt(16 if len(body_text) < 900 else 13)
                run.font.color.rgb = TEXT
                run.font.name = "Calibri"
                if para_line.startswith("    ") or (para_line and not para_line[:1].isalnum() and "-->" in para_line):
                    run.font.name = "Consolas"
                    run.font.size = PPt(12)
                    run.font.color.rgb = MUTED

        # footer
        add_textbox(
            slide,
            PInches(0.6),
            PInches(7.05),
            PInches(12),
            PInches(0.3),
            "Marvin Sim Chair · AVL Summer Cohort 2026",
            size=10,
            color=MUTED,
        )

    pptx_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(pptx_path))
    print("PPTX", pptx_path.name, "slides=", len(prs.slides))


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    docs = [
        "TECHNICAL_REPORT.md",
        "HANDOFF_PACKAGE.md",
        "DEMO_VIDEO_SCRIPT.md",
        "README.md",
    ]
    for name in docs:
        src = HERE / name
        if src.exists():
            md_to_docx(src, OUT / (src.stem + ".docx"))

    # deck section briefs too
    deck_dir = HERE.parent / "deck"
    if deck_dir.is_dir():
        for src in sorted(deck_dir.glob("*.md")):
            md_to_docx(src, OUT / (src.stem + ".docx"))

    slides = HERE / "SLIDE_DECK.md"
    if slides.exists():
        md_slides_to_pptx(slides, OUT / "SLIDE_DECK.pptx")

    print("Done. Outputs in", OUT)


if __name__ == "__main__":
    main()
